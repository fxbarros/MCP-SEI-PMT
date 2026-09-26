"""Modelos de parecer/despacho/ofício e geração de minutas.

Modelos ficam em iCloud Drive → "Modelos parecer" (sincronizado).
Minutas são salvas na MESMA pasta do processo (junto com o PDF/ZIP).
"""

from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Emu

from .sei_client import DESTINO_DOWNLOAD

PASTA_MODELOS = (
    Path.home()
    / "Library"
    / "Mobile Documents"
    / "com~apple~CloudDocs"
    / "Modelos parecer"
)
EXTENSOES_VALIDAS = (".docx", ".md", ".txt")


def _inferir_tipo(nome_lower: str) -> str:
    if "despacho" in nome_lower:
        return "despacho"
    if "ofício" in nome_lower or "oficio" in nome_lower:
        return "oficio"
    if "estrutura" in nome_lower:
        return "estrutura"
    if "parecer" in nome_lower:
        return "parecer"
    return "outro"


def listar_modelos() -> list[dict[str, Any]]:
    """Lista os arquivos de modelo disponíveis em iCloud / Modelos parecer/."""
    if not PASTA_MODELOS.exists():
        return []
    out: list[dict[str, Any]] = []
    for f in sorted(PASTA_MODELOS.iterdir()):
        if not f.is_file():
            continue
        # "~$nome.docx" = lock temporário do Word com o arquivo aberto; "." = oculto
        if f.name.startswith(("~$", ".")):
            continue
        if f.suffix.lower() not in EXTENSOES_VALIDAS:
            continue
        out.append(
            {
                "arquivo": f.name,
                "stem": f.stem,
                "extensao": f.suffix,
                "tipo_inferido": _inferir_tipo(f.stem.lower()),
                "tamanho_bytes": f.stat().st_size,
            }
        )
    return out


def ler_modelo(arquivo: str) -> dict[str, Any]:
    """Lê o texto de um modelo. Aceita nome com ou sem extensão."""
    if not PASTA_MODELOS.exists():
        raise FileNotFoundError(
            f"pasta de modelos não existe: {PASTA_MODELOS}"
        )

    path = PASTA_MODELOS / arquivo
    if not path.exists():
        # tenta match case-insensitive por nome ou stem
        alvo = arquivo.lower().strip()
        match: Path | None = None
        for f in PASTA_MODELOS.iterdir():
            if not f.is_file():
                continue
            if f.name.lower() == alvo or f.stem.lower() == alvo:
                match = f
                break
        if not match:
            raise FileNotFoundError(
                f"modelo {arquivo!r} não encontrado em {PASTA_MODELOS}. "
                f"Use listar_modelos_parecer() pra ver os disponíveis."
            )
        path = match

    if path.suffix.lower() == ".docx":
        doc = Document(str(path))
        texto = "\n".join(p.text for p in doc.paragraphs)
    else:
        texto = path.read_text(encoding="utf-8")

    return {
        "arquivo": path.name,
        "stem": path.stem,
        "extensao": path.suffix,
        "tipo_inferido": _inferir_tipo(path.stem.lower()),
        "texto": texto,
    }


# ---------------------------------------------------------------------------
# Heurísticas de classificação de parágrafos (usadas quando `modelo` é passado)
# ---------------------------------------------------------------------------
# Cabeçalho do parecer: linhas iniciais antes da EMENTA
_RE_CABECALHO = re.compile(
    r"^(Processo nº|Consulente:|Assunto:|Parecer Jurídico nº|Despacho nº|Ofício nº|Senhor[a]?,?$)",
    re.IGNORECASE,
)
# EMENTA (recebe recuo à esquerda como citação)
_RE_EMENTA = re.compile(r"^EMENTA[:\s]")
# Header de seção romano: "I. DO RELATÓRIO", "II. DA FUNDAMENTAÇÃO", etc.
_RE_HEADER_ROMANO = re.compile(r"^[IVX]+\.\s+[A-ZÁÉÍÓÚÂÊÔÃÕÇ ]{3,}$")
# Subheader: "II.1. Da finalidade ...", "II.4.1. ..."
_RE_SUBHEADER = re.compile(r"^[IVX]+\.\d+(?:\.\d+)?\.\s+\S")
# Parágrafo numerado: "1. ", "12. ", "29. "
_RE_NUMERADO = re.compile(r"^\d+\.\s+\S")
# Citação de lei / dispositivos transcritos
_RE_CITACAO = re.compile(
    r"^(Art\.\s|§\s?\d|Parágrafo\s|"
    r"[IVX]+\s*[\-—–]\s|"  # incisos: "I -", "II —", "III –"
    r"[IVX]+\s*[\-—–]?\s*[a-z]|"  # incisos minúsculos
    r"\[\.\.\.\]|"
    r"[a-z]\)\s)",  # alíneas: "a) "
)
# Marcadores explícitos de assinatura (centralizada)
_RE_ASSINATURA = re.compile(r"^(Teresina-?PI|Atenciosamente|Respeitosamente)", re.IGNORECASE)


def _extrair_perfis(template_path: Path) -> dict[str, dict[str, Any]]:
    """Lê o template e extrai os formatos típicos de cada tipo de parágrafo.

    Procura no template um exemplo de cada perfil (cabeçalho, ementa, header,
    numerado, citação, centralizado) e devolve seus parâmetros de formatação
    para que possamos replicá-los no novo documento.
    """
    perfis: dict[str, dict[str, Any]] = {}
    doc = Document(str(template_path))

    def snap(p) -> dict[str, Any]:
        return {
            "style": p.style.name,
            "alignment": p.alignment,
            "left_indent": p.paragraph_format.left_indent,
            "first_line_indent": p.paragraph_format.first_line_indent,
        }

    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            if "blank" not in perfis:
                perfis["blank"] = snap(p)
            continue
        if _RE_EMENTA.match(t) and "ementa" not in perfis:
            perfis["ementa"] = snap(p)
        elif _RE_HEADER_ROMANO.match(t) and "header" not in perfis:
            perfis["header"] = snap(p)
        elif _RE_SUBHEADER.match(t) and "subheader" not in perfis:
            perfis["subheader"] = snap(p)
        elif _RE_NUMERADO.match(t) and "numerado" not in perfis:
            perfis["numerado"] = snap(p)
        elif _RE_CITACAO.match(t) and "citacao" not in perfis:
            perfis["citacao"] = snap(p)
        elif p.alignment == WD_ALIGN_PARAGRAPH.CENTER and "centralizado" not in perfis:
            perfis["centralizado"] = snap(p)
        elif _RE_CABECALHO.match(t) and "cabecalho" not in perfis:
            perfis["cabecalho"] = snap(p)

    # Defaults caso o template não contenha algum perfil
    perfis.setdefault("blank", {"style": "Normal", "alignment": None, "left_indent": None, "first_line_indent": None})
    perfis.setdefault("cabecalho", {"style": "Normal", "alignment": WD_ALIGN_PARAGRAPH.JUSTIFY, "left_indent": None, "first_line_indent": None})
    perfis.setdefault("ementa", {"style": "Normal", "alignment": WD_ALIGN_PARAGRAPH.JUSTIFY, "left_indent": Emu(1524000), "first_line_indent": None})
    perfis.setdefault("header", {"style": "Normal", "alignment": WD_ALIGN_PARAGRAPH.JUSTIFY, "left_indent": None, "first_line_indent": None})
    perfis.setdefault("subheader", perfis["header"])
    perfis.setdefault("numerado", {"style": "Normal", "alignment": WD_ALIGN_PARAGRAPH.JUSTIFY, "left_indent": None, "first_line_indent": Emu(449580)})
    perfis.setdefault("citacao", perfis["ementa"])
    perfis.setdefault("centralizado", {"style": "Normal", "alignment": WD_ALIGN_PARAGRAPH.CENTER, "left_indent": None, "first_line_indent": None})
    return perfis


def _classificar(linha: str, modo_assinatura: bool) -> str:
    t = linha.strip()
    if not t:
        return "blank"
    if modo_assinatura:
        return "centralizado"
    if _RE_EMENTA.match(t):
        return "ementa"
    if _RE_HEADER_ROMANO.match(t):
        return "header"
    if _RE_SUBHEADER.match(t):
        return "subheader"
    if _RE_NUMERADO.match(t):
        return "numerado"
    if _RE_CITACAO.match(t):
        return "citacao"
    if _RE_CABECALHO.match(t):
        return "cabecalho"
    return "numerado"  # default razoável dentro do corpo do parecer


def _gerar_docx_com_modelo(
    template_path: Path, conteudo: str, destino: Path
) -> None:
    """Clona o template, esvazia o body e re-injeta o conteúdo classificando
    cada linha pelo perfil de formatação correspondente do próprio template."""
    perfis = _extrair_perfis(template_path)
    shutil.copy(template_path, destino)
    doc = Document(str(destino))

    body = doc.element.body
    sectPr_tag = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr"
    sectPr = body.find(f".//{sectPr_tag}")
    for child in list(body):
        if child.tag.endswith("}p") or child.tag.endswith("}tbl"):
            body.remove(child)

    modo_assinatura = False
    for linha in conteudo.split("\n"):
        if _RE_ASSINATURA.match(linha.strip()):
            modo_assinatura = True
        tipo_par = _classificar(linha, modo_assinatura)
        perfil = perfis[tipo_par]

        p = doc.add_paragraph()
        try:
            p.style = doc.styles[perfil["style"]]
        except KeyError:
            pass
        if perfil["alignment"] is not None:
            p.alignment = perfil["alignment"]
        if perfil["left_indent"] is not None:
            p.paragraph_format.left_indent = perfil["left_indent"]
        if perfil["first_line_indent"] is not None:
            p.paragraph_format.first_line_indent = perfil["first_line_indent"]

        if tipo_par == "blank":
            p.add_run("\xa0")
        else:
            p.add_run(linha)

    if sectPr is not None:
        body.append(sectPr)

    doc.core_properties.author = "Fábio Ximenes Barros (gerado por mcp-sei)"
    doc.core_properties.created = datetime.now()
    doc.save(str(destino))


def salvar_minuta(
    numero: str,
    conteudo: str,
    tipo: str = "Parecer",
    formato: str = "docx",
    modelo: str | None = None,
) -> Path:
    """Salva a minuta gerada na pasta do processo (junto ao PDF/ZIP).

    - `tipo` é o prefixo do nome ("Parecer", "Despacho", "Ofício", "Minuta", etc.)
    - `formato` em ('docx', 'md', 'txt')
    - `modelo`: nome do arquivo do template em "Modelos parecer" (com ou sem
      extensão). Quando passado e formato='docx', o template é clonado
      (preservando estilos, fontes, page setup) e o `conteudo` é re-injetado
      com classificação automática por tipo de parágrafo.
    - Caminho final: `~/Library/Mobile Documents/com~apple~CloudDocs/Processos SEI/{numero}/{tipo} {numero}.{formato}`
    - Se o arquivo já existe, é sobrescrito (cada análise gera versão fresca).
    """
    if formato not in ("docx", "md", "txt"):
        raise ValueError(f"formato {formato!r} inválido (use docx/md/txt)")

    numero_safe = numero.replace("/", "-")
    pasta = DESTINO_DOWNLOAD / numero_safe
    pasta.mkdir(parents=True, exist_ok=True)

    nome_arquivo = f"{tipo} {numero_safe}.{formato}"
    path = pasta / nome_arquivo

    if formato == "docx":
        if modelo:
            template = ler_modelo(modelo)  # valida existência e resolve nome
            template_path = PASTA_MODELOS / template["arquivo"]
            _gerar_docx_com_modelo(template_path, conteudo, path)
        else:
            doc = Document()
            doc.core_properties.author = "Fábio Ximenes Barros (gerado por mcp-sei)"
            doc.core_properties.created = datetime.now()
            for paragrafo in conteudo.split("\n"):
                doc.add_paragraph(paragrafo)
            doc.save(str(path))
    else:
        path.write_text(conteudo, encoding="utf-8")

    return path
